"""报告导出：PDF（WeasyPrint）与 Excel（OpenPyXL），以及采购需求汇总。

所有指标直接取方案落库的 metrics / solver_params 快照，保证"报告-界面-求解器"三处口径一致。
采购建议仅汇总需求（菜品配方 × 排餐份数 → 食材市品重量），不做库存。
"""
from __future__ import annotations

import io
import json
from collections import defaultdict
from datetime import datetime

from sqlmodel import Session, select
from weasyprint import HTML
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import MealPlan, MealItem, Elder, Dish, RecipeItem, Ingredient, AuditLog
from ..config import SLOT_LABELS, IDDSI_LABELS

NUTRIENT_LABELS = {
    "energy_kcal": ("能量", "kcal"),
    "protein_g": ("蛋白质", "g"),
    "fat_g": ("脂肪", "g"),
    "carbs_g": ("碳水化合物", "g"),
    "dietary_fiber_g": ("膳食纤维", "g"),
    "sodium_mg": ("钠", "mg"),
    "potassium_mg": ("钾", "mg"),
    "phosphorus_mg": ("磷", "mg"),
    "calcium_mg": ("钙", "mg"),
    "cholesterol_mg": ("胆固醇", "mg"),
    "sugar_g": ("糖", "g"),
    "purine_mg": ("嘌呤", "mg"),
}
RADAR_KEYS = ["energy_kcal", "protein_g", "dietary_fiber_g", "calcium_mg",
              "potassium_mg", "phosphorus_mg", "sodium_mg"]


def load_plan_bundle(session: Session, plan_id: int) -> dict:
    plan = session.get(MealPlan, plan_id)
    if not plan:
        raise ValueError("方案不存在")
    elder = session.get(Elder, plan.elder_id)
    items = session.exec(select(MealItem).where(MealItem.plan_id == plan_id)
                         .order_by(MealItem.day_index, MealItem.slot)).all()
    dishes = {d.id: d for d in session.exec(select(Dish)).all()}
    metrics = json.loads(plan.metrics_json or "{}")
    params = json.loads(plan.solver_params or "{}")
    target = {}
    try:
        target = json.loads(elder.target_explanation or "{}")
    except json.JSONDecodeError:
        pass
    return {"plan": plan, "elder": elder, "items": items, "dishes": dishes,
            "metrics": metrics, "params": params, "target": target}


# ───────────────── 采购需求汇总 ─────────────────

def purchase_summary(session: Session, plan_id: int) -> list[dict]:
    bundle = load_plan_bundle(session, plan_id)
    agg: dict[int, dict] = {}
    recipe_rows = session.exec(select(RecipeItem)).all()
    recipes: dict[int, list[RecipeItem]] = defaultdict(list)
    for r in recipe_rows:
        recipes[r.dish_id].append(r)
    ings = {i.id: i for i in session.exec(select(Ingredient)).all()}

    for it in bundle["items"]:
        dish = bundle["dishes"].get(it.dish_id)
        if not dish:
            continue
        scale = it.portion_g / max(dish.portion_g, 1)
        for r in recipes.get(dish.id, []):
            ing = ings.get(r.ingredient_id)
            if not ing:
                continue
            row = agg.setdefault(r.ingredient_id, {
                "ingredient": ing.name, "code": ing.code,
                "category": ing.category, "gross_g": 0.0,
                "edible_g": 0.0, "cost": 0.0})
            gross = r.gross_weight_g * scale
            row["gross_g"] += gross
            row["edible_g"] += gross * ing.edible_rate * (1 - r.cooking_loss_rate)
            row["cost"] += gross * ing.unit_cost / 100.0
    out = sorted(agg.values(), key=lambda x: (x["category"], -x["gross_g"]))
    for r in out:
        r["gross_g"] = round(r["gross_g"], 0)
        r["edible_g"] = round(r["edible_g"], 0)
        r["cost"] = round(r["cost"], 2)
    return out


# ───────────────── PDF ─────────────────

CSS = """
@page { size: A4; margin: 1.6cm 1.4cm; @bottom-center { content: "第 " counter(page) " 页 / 共 " counter(pages) " 页"; font-size: 9px; color:#888; } }
body { font-family: "Noto Sans CJK SC","Noto Sans CJK SC Regular",sans-serif; font-size: 11px; color:#222; }
h1 { font-size: 20px; margin: 0 0 4px; color:#1f3a5f; }
h2 { font-size: 14px; margin: 18px 0 6px; color:#1f3a5f; border-left: 4px solid #2f6fb0; padding-left:8px; }
.meta { color:#555; font-size: 10px; margin-bottom: 10px; }
table { width: 100%; border-collapse: collapse; margin: 6px 0 12px; }
th, td { border: 1px solid #bbb; padding: 4px 6px; text-align: left; vertical-align: top; }
th { background: #e8f0fa; font-weight: 600; }
.day-title { background:#f2f7fd; padding:5px 8px; margin-top:10px; font-weight:600; color:#1f3a5f; }
.warn { color:#b45309; } .forbid { color:#b91c1c; font-weight:600; }
.ok { color:#15803d; }
.small { font-size: 9px; color:#666; }
.tag { display:inline-block; background:#eef; border-radius:3px; padding:0 4px; margin-right:4px; }
.kpi { display:inline-block; background:#f2f7fd; border:1px solid #cfe0f3; border-radius:6px;
       padding:6px 12px; margin:0 8px 8px 0; }
.kpi b { font-size:15px; color:#1f3a5f; }
ul { margin: 4px 0; padding-left: 18px; }
"""


def _nutrient_table_rows(bundle: dict) -> str:
    target = bundle["target"].get("targets", {})
    daily = bundle["metrics"].get("daily_nutrients", [])
    days = len(daily) or 1
    rows = ""
    for k, (label, unit) in NUTRIENT_LABELS.items():
        vals = [d.get(k, 0) for d in daily]
        avg = sum(vals) / days
        tgt = target.get(k)
        if tgt:
            cls = "ok" if tgt[0] <= avg <= tgt[1] else "warn"
            tgt_txt = f"{tgt[0]}~{tgt[1]}"
        else:
            cls, tgt_txt = "", "—"
        rows += (f"<tr><td>{label}</td><td>{avg:.1f}</td><td>{tgt_txt}</td>"
                 f"<td class='{cls}'>{'达标' if cls=='ok' else ('偏离' if cls=='warn' else '—')}</td>"
                 f"<td class='small'>{unit}</td></tr>")
    return rows


def _day_sections(bundle: dict) -> str:
    plan = bundle["plan"]
    days = 1
    if plan.period_type == "week":
        d0 = datetime.fromisoformat(plan.start_date)
        d1 = datetime.fromisoformat(plan.end_date)
        days = (d1 - d0).days + 1
    html = ""
    for d in range(days):
        html += f"<div class='day-title'>第 {d+1} 天</div><table><tr><th style='width:70px'>餐次</th><th>菜品</th><th style='width:70px'>份量(g)</th><th style='width:80px'>成本(元)</th><th style='width:60px'>质地</th></tr>"
        day_items = [i for i in bundle["items"] if i.day_index == d]
        for slot in ("breakfast", "lunch", "dinner"):
            slot_items = [i for i in day_items if i.slot == slot]
            for j, it in enumerate(slot_items):
                dish = bundle["dishes"].get(it.dish_id)
                detail = json.loads(it.detail_json or "{}")
                lock = " 🔒" if it.locked else ""
                html += ("<tr>" +
                         (f"<td rowspan='{len(slot_items)}'>{SLOT_LABELS.get(slot, slot)}</td>" if j == 0 else "") +
                         f"<td>{it.dish_name}{lock}</td><td>{it.portion_g:.0f}</td>"
                         f"<td>{detail.get('cost', 0)}</td>"
                         f"<td>{dish.iddsi_level if dish else ''}级</td></tr>")
        dc = bundle["metrics"].get("daily_costs", [])
        day_cost = dc[d] if d < len(dc) else 0
        html += f"<tr><td colspan='3' style='text-align:right'>当日合计</td><td colspan='2'><b>{day_cost}</b> 元</td></tr>"
        html += "</table>"
    return html


def render_plan_pdf(session: Session, plan_id: int) -> bytes:
    b = load_plan_bundle(session, plan_id)
    p, e, m = b["plan"], b["elder"], b["metrics"]
    bd = m.get("score_breakdown", {})
    warns = m.get("warnings", [])
    relax = m.get("relaxations", [])
    explain = b["target"].get("explanation", [])
    conflicts_notes = b["target"].get("conflict_notes", [])
    kpis = f"""
    <div><span class='kpi'>日均成本 <b>{bd.get('avg_day_cost','—')}</b> 元</span>
    <span class='kpi'>总成本 <b>{m.get('total_cost','—')}</b> 元</span>
    <span class='kpi'>平均营养偏差 <b>{bd.get('avg_nutrient_deviation_pct','—')}%</b></span>
    <span class='kpi'>满意度 <b>{bd.get('avg_satisfaction','—')}</b>/5</span>
    <span class='kpi'>浪费估算 <b>{bd.get('waste_g_est','—')}</b> g</span></div>"""
    notes = ""
    if relax:
        notes += "<h2>求解松弛记录</h2><ul>" + "".join(f"<li class='warn'>{x}</li>" for x in relax) + "</ul>"
    if warns:
        notes += "<h2>求解警告</h2><ul>" + "".join(f"<li class='warn'>{x}</li>" for x in warns) + "</ul>"
    if conflicts_notes:
        notes += "<h2>多病共存冲突调解说明</h2><ul>" + "".join(f"<li class='forbid'>{x}</li>" for x in conflicts_notes) + "</ul>"
    html = f"""<!doctype html><html><head><meta charset='utf-8'><style>{CSS}</style></head><body>
    <h1>老年营养配餐方案</h1>
    <div class='meta'>
      方案：{p.title} ｜ 状态：{p.status} ｜ 周期：{p.start_date} ~ {p.end_date or p.start_date} ｜ 规则版本：{p.rule_version}<br/>
      老人：{e.name} ｜ 性别：{'男' if e.gender=='male' else '女'} ｜ 身高 {e.height_cm}cm ｜ 体重 {e.weight_kg}kg
      ｜ 吞咽：{IDDSI_LABELS.get(e.iddsi_level, e.iddsi_level)} ｜ 宗教：{e.religion}
      ｜ 慢病：{e.chronic_diseases or '无'} ｜ 过敏：{e.allergies or '无'} ｜ 药物：{e.medications or '无'}<br/>
      生成时间：{p.created_at:%Y-%m-%d %H:%M} ｜ 创建人：{p.created_by} ｜ 求解状态：{bd.get('status', m.get('daily_nutrients') and 'ok')}
    </div>
    <h2>核心指标</h2>{kpis}
    <h2>排餐明细</h2>{_day_sections(b)}
    <h2>每日营养摄入与目标对比（多日平均）</h2>
    <table><tr><th>营养素</th><th>实际均值</th><th>目标区间</th><th>判定</th><th>单位</th></tr>
    {_nutrient_table_rows(b)}</table>
    {notes}
    <h2>个体化目标计算依据</h2><ul class='small'>
    {''.join(f'<li>{x}</li>' for x in explain)}</ul>
    <h2>求解器参数快照</h2>
    <pre class='small'>{json.dumps({k:v for k,v in b['params'].items() if k!='targets'}, ensure_ascii=False, indent=2)}</pre>
    </body></html>"""
    return HTML(string=html).write_pdf()


# ───────────────── Excel ─────────────────

THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DCE9F8")
HEAD_FONT = Font(bold=True)


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _autofit(ws, max_width=40):
    for col in ws.columns:
        width = 8
        letter = get_column_letter(col[0].column)
        for c in col:
            if c.value is not None:
                width = max(width, min(max_width, len(str(c.value)) + 2))
        ws.column_dimensions[letter].width = width


def render_plan_excel(session: Session, plan_id: int) -> bytes:
    b = load_plan_bundle(session, plan_id)
    p, e, m = b["plan"], b["elder"], b["metrics"]
    wb = Workbook()

    ws = wb.active
    ws.title = "方案概览"
    ws.append(["老年营养配餐方案"])
    ws["A1"].font = Font(bold=True, size=14)
    overview = [
        ("方案标题", p.title), ("状态", p.status), ("周期", f"{p.start_date} ~ {p.end_date or p.start_date}"),
        ("规则版本", p.rule_version), ("老人", e.name), ("性别", "男" if e.gender == "male" else "女"),
        ("身高(cm)", e.height_cm), ("体重(kg)", e.weight_kg),
        ("吞咽等级", IDDSI_LABELS.get(e.iddsi_level, str(e.iddsi_level))),
        ("慢病", e.chronic_diseases), ("过敏", e.allergies), ("宗教禁忌", e.religion),
        ("药物", e.medications), ("忌口", e.dislikes), ("营养目标", e.nutrition_goal),
        ("日均成本上限", e.cost_limit_day), ("实际日均成本",
            m.get("score_breakdown", {}).get("avg_day_cost")),
        ("平均营养偏差%", m.get("score_breakdown", {}).get("avg_nutrient_deviation_pct")),
        ("满意度", m.get("score_breakdown", {}).get("avg_satisfaction")),
        ("浪费估算(g)", m.get("score_breakdown", {}).get("waste_g_est")),
        ("求解状态", m.get("status")), ("求解耗时(s)", m.get("solve_seconds")),
    ]
    for k, v in overview:
        ws.append([k, v])
    _autofit(ws)

    ws2 = wb.create_sheet("排餐明细")
    ws2.append(["天", "日期", "餐次", "菜品", "菜品类型", "份量(g)", "成本(元)", "IDDSI", "锁定"])
    _style_header(ws2)
    for it in b["items"]:
        dish = b["dishes"].get(it.dish_id)
        detail = json.loads(it.detail_json or "{}")
        ws2.append([it.day_index + 1, it.meal_date, SLOT_LABELS.get(it.slot, it.slot),
                    it.dish_name, detail.get("dish_type", ""), it.portion_g,
                    detail.get("cost", 0), dish.iddsi_level if dish else "",
                    "是" if it.locked else ""])
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
    _autofit(ws2)

    ws3 = wb.create_sheet("每日营养")
    target = b["target"].get("targets", {})
    header = ["天"] + [NUTRIENT_LABELS[k][0] for k in NUTRIENT_LABELS] + ["当日成本(元)"]
    ws3.append(header)
    _style_header(ws3)
    for idx, dn in enumerate(m.get("daily_nutrients", [])):
        ws3.append([idx + 1] + [dn.get(k, 0) for k in NUTRIENT_LABELS]
                   + [m.get("daily_costs", [0] * 99)[idx] if idx < len(m.get("daily_costs", [])) else 0])
    ws3.append(["目标区间"] + [
        f"{target[k][0]}~{target[k][1]}" if k in target else "—"
        for k in NUTRIENT_LABELS] + [f"≤{e.cost_limit_day}"])
    for c in ws3[ws3.max_row]:
        c.font = Font(italic=True, color="1f3a5f")
    _autofit(ws3)

    ws4 = wb.create_sheet("采购需求汇总")
    ws4.append(["分类", "食材", "编码", "市品需求量(g)", "折可食熟重(g)", "估算成本(元)"])
    _style_header(ws4)
    for r in purchase_summary(session, plan_id):
        ws4.append([r["category"], r["ingredient"], r["code"],
                    r["gross_g"], r["edible_g"], r["cost"]])
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
    _autofit(ws4)

    ws5 = wb.create_sheet("求解与规则记录")
    ws5.append(["项目", "内容"])
    _style_header(ws5)
    params = b["params"]
    for k, v in params.items():
        ws5.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
    for note in m.get("relaxations", []):
        ws5.append(["松弛记录", note])
    for note in m.get("warnings", []):
        ws5.append(["警告", note])
    _autofit(ws5)

    ws6 = wb.create_sheet("操作日志")
    ws6.append(["时间", "操作人", "动作", "对象", "详情"])
    _style_header(ws6)
    logs = session.exec(select(AuditLog)
                        .where(AuditLog.entity_id == str(plan_id))
                        .order_by(AuditLog.id)).all()
    for lg in logs:
        ws6.append([lg.created_at.strftime("%Y-%m-%d %H:%M:%S"), lg.username,
                    lg.action, lg.entity_type,
                    json.dumps(json.loads(lg.detail_json or "{}"), ensure_ascii=False)])
    _autofit(ws6)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_compare_excel(session: Session, plan_ids: list[int]) -> bytes:
    """多方案对比 Excel。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "方案对比"
    bundles = [load_plan_bundle(session, pid) for pid in plan_ids]
    ws.append(["指标"] + [b["plan"].title for b in bundles])
    _style_header(ws)
    metric_rows = [
        ("状态", lambda b: b["plan"].status),
        ("周期", lambda b: f"{b['plan'].start_date}~{b['plan'].end_date}"),
        ("规则版本", lambda b: b["plan"].rule_version),
        ("总成本(元)", lambda b: b["metrics"].get("total_cost")),
        ("日均成本(元)", lambda b: b["metrics"].get("score_breakdown", {}).get("avg_day_cost")),
        ("平均营养偏差%", lambda b: b["metrics"].get("score_breakdown", {})
            .get("avg_nutrient_deviation_pct")),
        ("满意度", lambda b: b["metrics"].get("score_breakdown", {}).get("avg_satisfaction")),
        ("浪费估算(g)", lambda b: b["metrics"].get("score_breakdown", {}).get("waste_g_est")),
        ("重复附加次数", lambda b: b["metrics"].get("score_breakdown", {}).get("repeat_extra_count")),
        ("软风险命中", lambda b: b["metrics"].get("score_breakdown", {}).get("risk_tag_hits")),
        ("不同菜品数", lambda b: b["metrics"].get("score_breakdown", {}).get("distinct_dishes")),
        ("求解耗时(s)", lambda b: b["metrics"].get("solve_seconds")),
    ]
    for label, fn in metric_rows:
        ws.append([label] + [fn(b) for b in bundles])

    ws2 = wb.create_sheet("每日营养均值")
    target = bundles[0]["target"].get("targets", {})
    ws2.append(["营养素", "单位", "目标区间"] + [b["plan"].title for b in bundles])
    _style_header(ws2)
    days_of = lambda b: max(len(b["metrics"].get("daily_nutrients", [])), 1)
    for k, (label, unit) in NUTRIENT_LABELS.items():
        row = [label, unit,
               f"{target[k][0]}~{target[k][1]}" if k in target else "—"]
        for b in bundles:
            dn = b["metrics"].get("daily_nutrients", [])
            avg = sum(d.get(k, 0) for d in dn) / days_of(b) if dn else 0
            row.append(round(avg, 1))
        ws2.append(row)
    _autofit(ws)
    _autofit(ws2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
