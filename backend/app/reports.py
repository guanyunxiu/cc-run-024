"""报告导出：营养方案 PDF（WeasyPrint）、Excel（OpenPyXL）、采购需求汇总。

采购清单仅做需求汇总（按食材合计毛重），不涉及仓库库存。
"""
import io
from collections import defaultdict
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlmodel import Session, select

from .config import MEAL_SLOTS
from .models import Dish, DishIngredient, Ingredient, MenuItem, MenuPlan, Resident
from .solver import NUT_LABELS, NUTRIENTS, SLOT_LINES

SLOT_LABELS = {k: zh for k, zh, _ in MEAL_SLOTS}
LINE_LABELS = {"staple": "主食", "protein": "蛋白菜", "entree": "主菜",
               "vegetable": "蔬菜", "soup": "汤", "snack": "加餐"}


# ---------------- 数据组装 ----------------

def build_report(session: Session, plan: MenuPlan) -> dict:
    resident = session.get(Resident, plan.resident_id)
    items = session.exec(select(MenuItem).where(
        MenuItem.plan_id == plan.id)).all()
    dish_map = {d.id: d for d in session.exec(select(Dish)).all()}

    # 日 × 餐次网格
    grid = {}
    for it in items:
        grid.setdefault(it.day_index, {}).setdefault(it.slot, []).append(it)

    dates = [(plan.start_date + timedelta(days=d)).strftime("%Y-%m-%d")
             for d in range(plan.days)]
    day_totals = (plan.totals or {}).get("day_totals", {})
    return {
        "plan": plan, "resident": resident, "items": items,
        "dish_map": dish_map, "grid": grid, "dates": dates,
        "day_totals": day_totals,
        "score": plan.score or {}, "totals": plan.totals or {},
        "targets": (plan.targets_snapshot or {}).get("targets", {}),
        "assumptions": (plan.targets_snapshot or {}).get("assumptions", []),
        "conflicts_calc": (plan.targets_snapshot or {}).get("conflicts", []),
        "rule_version": plan.rule_version,
        "solver_params": plan.solver_params or {},
        "soft_violations": (plan.totals or {}).get("soft_violations", []),
        "warnings": (plan.totals or {}).get("warnings", []),
    }


def procurement_list(session: Session, plan: MenuPlan) -> dict:
    """按整个方案周期（默认每餐 1 份）汇总食材毛重需求与预估成本。"""
    items = session.exec(select(MenuItem).where(
        MenuItem.plan_id == plan.id)).all()
    recipe_rows = session.exec(select(DishIngredient)).all()
    recipe_by_dish: dict[int, list[DishIngredient]] = defaultdict(list)
    for r in recipe_rows:
        recipe_by_dish[r.dish_id].append(r)
    ing_map = {i.id: i for i in session.exec(select(Ingredient)).all()}

    agg = defaultdict(lambda: {"gross_g": 0.0, "servings": 0})
    for it in items:
        for r in recipe_by_dish.get(it.dish_id, []):
            agg[r.ingredient_id]["gross_g"] += r.gross_g
            agg[r.ingredient_id]["servings"] += 1

    rows = []
    total_cost = 0.0
    cat_group = defaultdict(list)
    for ing_id, v in agg.items():
        ing = ing_map.get(ing_id)
        if not ing:
            continue
        cost = v["gross_g"] / 1000.0 * ing.unit_cost
        total_cost += cost
        row = {"ingredient_id": ing_id, "name": ing.name,
               "category": ing.category, "gross_g": round(v["gross_g"], 1),
               "gross_kg": round(v["gross_g"] / 1000, 2),
               "unit_cost": ing.unit_cost, "cost": round(cost, 2),
               "edible_pct": ing.edible_pct,
               "edible_weight_kg": round(v["gross_g"] * ing.edible_pct / 100 / 1000, 2)}
        rows.append(row)
        cat_group[ing.category].append(row)
    rows.sort(key=lambda x: (x["category"], -x["gross_g"]))
    return {"rows": rows, "total_cost": round(total_cost, 2),
            "days": plan.days, "categories": sorted(cat_group)}


# ---------------- Excel ----------------

def export_excel(session: Session, plan: MenuPlan) -> bytes:
    rep = build_report(session, plan)
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="2F5496")
    hdr_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1) 方案总览
    ws = wb.active
    ws.title = "方案总览"
    r = rep
    info = [
        ("方案名称", r["plan"].name), ("老人", r["resident"].name),
        ("起始日期", str(r["plan"].start_date)), ("天数", r["plan"].days),
        ("状态", {"draft": "草稿", "published": "已发布",
                  "archived": "已归档"}.get(r["plan"].status, r["plan"].status)),
        ("日预算(元)", r["plan"].budget_per_day),
        ("实际日均成本(元)", r["score"].get("avg_cost_per_day")),
        ("平均营养达标率(%)", r["score"].get("avg_attainment_pct")),
        ("平均满意度", r["score"].get("avg_satisfaction")),
        ("菜品丰富度", r["score"].get("distinct_dishes")),
        ("软约束提示数", r["score"].get("soft_violation_count")),
        ("规则版本", r["rule_version"]),
        ("求解器参数", str(r["solver_params"])),
    ]
    ws.append(["养老机构营养配餐方案"])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in info:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 45

    # 2) 每日配餐
    ws2 = wb.create_sheet("每日配餐")
    ws2.append(["天数", "日期", "餐次", "线位", "菜品", "IDDSI", "成本(元)", "锁定"])
    for c in ws2[1]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, center
    for day in range(r["plan"].days):
        for slot, _zh, _pct in MEAL_SLOTS:
            for it in sorted(r["grid"].get(day, {}).get(slot, []),
                             key=lambda x: x.line):
                d = r["dish_map"].get(it.dish_id)
                ws2.append([f"第{day+1}天", r["dates"][day], SLOT_LABELS[slot],
                            LINE_LABELS.get(it.line, it.line), d.name if d else "",
                            d.iddsi_level if d else "", round(d.cost, 2) if d else 0,
                            "是" if it.locked else ""])
    for col, wdt in zip("ABCDEFGH", [8, 12, 10, 8, 24, 8, 10, 6]):
        ws2.column_dimensions[col].width = wdt

    # 3) 每日营养
    ws3 = wb.create_sheet("每日营养")
    header = ["天数/指标"] + [NUT_LABELS[n] for n in NUTRIENTS] + ["成本(元)"]
    ws3.append(header)
    for c in ws3[1]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, center
    for day in range(r["plan"].days):
        dt = r["day_totals"].get(str(day), {})
        ws3.append([f"第{day+1}天"] + [round(dt.get(n, 0), 1) for n in NUTRIENTS]
                   + [round(dt.get("cost", 0), 2)])
    # 目标行
    t = r["targets"]
    def tv(n, k):
        return (t.get(n) or {}).get(k, "")
    ws3.append(["目标下限"] + [tv(n, "min") for n in NUTRIENTS] + [""])
    ws3.append(["目标上限"] + [tv(n, "max") for n in NUTRIENTS]
               + [r["plan"].budget_per_day])
    ws3.append(["达标率%"] + [r["score"].get("attainment", {}).get(n, "")
                             for n in NUTRIENTS] + [""])
    for i, col in enumerate("ABCDEFGHIJK"):
        ws3.column_dimensions[col].width = 12 if i == 0 else 11

    # 4) 软约束提示
    ws4 = wb.create_sheet("软约束提示")
    ws4.append(["天数", "餐次", "菜品", "规则", "级别", "说明", "依据"])
    for c in ws4[1]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, center
    for v in r["soft_violations"]:
        ws4.append([f"第{v.get('day_index',0)+1}天", v.get("slot"),
                    v.get("dish_name"), v.get("rule_name"), "软约束",
                    v.get("message"), v.get("rationale")])
    for w_ in r["warnings"]:
        ws4.append(["-", "-", "-", "求解器提示", "警告", w_, ""])
    for col, wdt in zip("ABCDEFG", [8, 10, 18, 16, 8, 45, 40]):
        ws4.column_dimensions[col].width = wdt

    # 5) 采购需求汇总
    ws5 = wb.create_sheet("采购需求汇总")
    proc = procurement_list(session, plan)
    ws5.append(["食材分类", "食材", "需求毛重(kg)", "可食部%",
                "可食量(kg)", "单价(元/kg)", "预估金额(元)"])
    for c in ws5[1]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, center
    for row in proc["rows"]:
        ws5.append([row["category"], row["name"], row["gross_kg"],
                    row["edible_pct"], row["edible_weight_kg"],
                    row["unit_cost"], row["cost"]])
    ws5.append(["合计", "", "", "", "", "", proc["total_cost"]])
    ws5.append(["说明", "本清单为需求汇总，不含库存扣减", "", "", "", "", ""])
    for col, wdt in zip("ABCDEFG", [10, 16, 13, 9, 11, 12, 13]):
        ws5.column_dimensions[col].width = wdt

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------

_FONT_PATH = "/workspace/backend/assets/fonts/SimHei.ttf"


def _esc(s) -> str:
    return (str(s if s is not None else "")
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def export_pdf(session: Session, plan: MenuPlan) -> bytes:
    from weasyprint import HTML

    rep = build_report(session, plan)
    r = rep
    # 配餐表
    meal_rows = []
    for day in range(r["plan"].days):
        cells = ""
        for slot, _zh, _pct in MEAL_SLOTS:
            dishes = []
            for it in sorted(r["grid"].get(day, {}).get(slot, []),
                             key=lambda x: x.line):
                d = r["dish_map"].get(it.dish_id)
                if d:
                    lock = " 🔒" if it.locked else ""
                    dishes.append(
                        f"{_esc(d.name)}"
                        f"<span class='meta'>（L{d.iddsi_level}/¥{d.cost:.1f}）{lock}</span>")
            cells += (f"<div class='meal'><b>{SLOT_LABELS[slot]}</b>："
                      f"{'、'.join(dishes) if dishes else '—'}</div>")
        meal_rows.append(
            f"<div class='daycard'><h3>第{day+1}天 · {r['dates'][day]}</h3>{cells}</div>")

    # 营养达标表
    att = r["score"].get("attainment", {})
    nut_rows = []
    for n in NUTRIENTS:
        spec = r["targets"].get(n) or {}
        avg = (r["day_totals"].get("0") or {})
        avg_val = None
        vals = [r["day_totals"].get(str(d), {}).get(n, 0)
                for d in range(r["plan"].days)]
        avg_val = sum(vals) / max(len(vals), 1)
        pct = att.get(n, "")
        color = "#1a7f37" if isinstance(pct, (int, float)) and pct >= 85 else "#b26a00"
        nut_rows.append(
            f"<tr><td>{NUT_LABELS[n]}</td><td>{spec.get('min','—')}</td>"
            f"<td>{spec.get('max','—')}</td><td>{avg_val:.1f}</td>"
            f"<td style='color:{color};font-weight:bold'>{pct}%</td></tr>")

    # 软约束
    soft_html = "".join(
        f"<li>第{v.get('day_index',0)+1}天 {SLOT_LABELS.get(v.get('slot'),'')}"
        f"《{_esc(v.get('dish_name'))}》：{_esc(v.get('message'))}"
        f"<div class='rat'>依据：{_esc(v.get('rationale'))}</div></li>"
        for v in r["soft_violations"][:30]) or "<li>无软约束提示</li>"
    warn_html = "".join(f"<li>{_esc(w)}</li>" for w in r["warnings"]) or "<li>无</li>"
    assumption_html = "".join(f"<li>{_esc(a)}</li>" for a in r["assumptions"]) or "<li>—</li>"
    conflict_html = "".join(
        f"<li><b>{_esc(c.get('topic'))}</b><br/>折衷：{_esc(c.get('resolution'))}</li>"
        for c in r["conflicts_calc"]) or "<li>无冲突</li>"

    s = r["score"]
    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
@font-face {{ font-family: SimHei; src: url('file://{_FONT_PATH}'); }}
* {{ font-family: SimHei, sans-serif; }}
body {{ font-size: 12px; color:#222; margin: 18px; }}
h1 {{ font-size: 20px; text-align:center; margin-bottom: 2px; }}
.sub {{ text-align:center; color:#666; margin-bottom:14px; }}
h2 {{ font-size:14px; border-left:4px solid #2f5496; padding-left:8px; margin-top:18px; }}
h3 {{ font-size:13px; margin:6px 0; }}
table {{ border-collapse:collapse; width:100%; margin-top:6px; }}
th,td {{ border:1px solid #bbb; padding:5px 7px; font-size:11px; }}
th {{ background:#2f5496; color:#fff; }}
.daycard {{ border:1px solid #ccc; border-radius:6px; padding:8px 10px; margin:8px 0; page-break-inside:avoid; }}
.meal {{ margin:3px 0; line-height:1.5; }}
.meta {{ color:#888; font-size:10px; }}
.kpis {{ display:flex; flex-wrap:wrap; gap:8px; margin:8px 0; }}
.kpi {{ border:1px solid #2f5496; border-radius:6px; padding:6px 10px; min-width:110px; }}
.kpi b {{ font-size:16px; color:#2f5496; }}
.rat {{ color:#777; font-size:10px; }}
ul {{ margin:4px 0; padding-left:20px; }} li {{ margin:3px 0; }}
.foot {{ margin-top:20px; color:#888; font-size:10px; text-align:center; }}
</style></head><body>
<h1>老年营养与慢病配餐方案</h1>
<div class="sub">{_esc(r['plan'].name)} ｜ 生成于 {__import__('datetime').datetime.now().strftime('%Y-%m-%d %H:%M')}</div>

<h2>一、基本信息</h2>
<table>
<tr><th>老人</th><td>{_esc(r['resident'].name)}</td><th>年龄/性别</th><td>{r['resident'].age}岁 / {'男' if r['resident'].gender=='male' else '女'}</td></tr>
<tr><th>身高体重</th><td>{r['resident'].height_cm}cm / {r['resident'].weight_kg}kg（BMI {r['plan'].targets_snapshot.get('bmi','—')}）</td><th>吞咽等级</th><td>IDDSI {r['resident'].swallowing_level}</td></tr>
<tr><th>慢病</th><td>{_esc('、'.join(r['resident'].chronic_conditions) or '无')}</td><th>过敏</th><td>{_esc('、'.join(r['resident'].allergies) or '无')}</td></tr>
<tr><th>宗教</th><td>{_esc(r['resident'].religion)}</td><th>用药</th><td>{_esc('、'.join(r['resident'].medications) or '无')}</td></tr>
<tr><th>周期</th><td>{_esc(r['plan'].start_date)} 起 {r['plan'].days} 天</td><th>日预算</th><td>¥{r['plan'].budget_per_day:.1f}</td></tr>
</table>

<div class="kpis">
  <div class="kpi">平均达标率<br/><b>{s.get('avg_attainment_pct','—')}%</b></div>
  <div class="kpi">日均成本<br/><b>¥{s.get('avg_cost_per_day','—')}</b></div>
  <div class="kpi">满意度<br/><b>{s.get('avg_satisfaction','—')}</b></div>
  <div class="kpi">菜品丰富度<br/><b>{s.get('distinct_dishes','—')}</b></div>
  <div class="kpi">软约束提示<br/><b>{s.get('soft_violation_count','—')}</b></div>
</div>

<h2>二、每日配餐</h2>
{''.join(meal_rows)}

<h2>三、营养目标与达标（日均值）</h2>
<table><tr><th>指标</th><th>下限</th><th>上限</th><th>实际日均</th><th>达标率</th></tr>
{''.join(nut_rows)}</table>

<h2>四、目标计算依据与多病共存冲突</h2>
<b>计算假设：</b><ul>{assumption_html}</ul>
<b>冲突与折衷：</b><ul>{conflict_html}</ul>

<h2>五、软约束与可解释提示</h2>
<ul>{soft_html}</ul>
<b>求解器告警：</b><ul>{warn_html}</ul>

<h2>六、审计信息</h2>
<table>
<tr><th>规则版本</th><td>{_esc(r['rule_version'])}</td></tr>
<tr><th>求解器参数</th><td>{_esc(str(r['solver_params']))}</td></tr>
<tr><th>求解状态/用时</th><td>{_esc(s.get('status'))} / {s.get('solve_time_sec','—')} 秒</td></tr>
<tr><th>创建人</th><td>{_esc(r['plan'].created_by)} ｜ {r['plan'].created_at}</td></tr>
</table>
<div class="foot">本报告由老年营养与慢病配餐优化管理系统自动生成，仅供营养专业人员参考，临床决策请结合医嘱。</div>
</body></html>"""

    return HTML(string=html, base_url="/workspace/backend").write_pdf()
